from openpyxl import Workbook
import pandas as pd
from openpyxl.chart import BarChart, Series, Reference


def make_bar_chart(ws, bar_title, x_title, y_title):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = bar_title
    chart.y_axis.title = y_title
    chart.x_axis.title = x_title
    data = Reference(ws, min_col=1, min_row=1, max_row=8, max_col=2)
    cats = Reference(ws, min_col=1, min_row=2, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.shape = 4
    ws.add_chart(chart, "A10")
    return ws


if __name__ == "__main__":
    row_data = {"日時":["4月", "5月", "6月", "7月", "8月", "9月", "10月"],
    "データ件数": [3, 4, 5, 6, 11, 3, 11]}
    df = pd.DataFrame(row_data)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ordered_column = ["日時", "データ件数"]

    # セルにデータ入力
    # 列を入力
    ws.append(ordered_column)
    # 行ごとに入力
    df.apply(lambda x: ws.append(x[ordered_column].values.tolist()), axis=1)
    ws = make_bar_chart(ws, "hoge", "xxx", "yyy")
    wb.save("hoge.xlsx")

